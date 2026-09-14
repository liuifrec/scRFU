"""Freeze evidence axes and locus dossiers from the completed real-data pilot.

Requires the verified full nominal/conditional lookup and versioned API output.
Does not infer RFU colocalization from gene/peak proximity or source GWAS coloc.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import subprocess
import sys
import zipfile
from pathlib import Path

import pandas as pd

try:
    from .matos_regulatory_lookup import digest
    from .run_regulatory_realdata_pilot import read, save
except ImportError:
    from matos_regulatory_lookup import digest
    from run_regulatory_realdata_pilot import read, save


# NCBI Gene 6957, GRCh38.p14 / NC_000007.14, one-based inclusive.
TRB = (142299011, 142813287)
SUPPLEMENT_MD5 = "2a8ed13f379946c1a7efce117c961201"


def location(chrom: str, position: int) -> str:
    if str(chrom) != "7":
        return "other"
    if TRB[0] <= position <= TRB[1]:
        return "TRB_locus"
    if TRB[0] - 250_000 <= position <= TRB[1] + 250_000:
        return "TRB_flanking"
    return "other"


def evidence_tier(eq_ind: bool, ca_ind: bool, eq_support: bool, ca_support: bool) -> str:
    if eq_ind and ca_ind:
        return "A_same_variant_conditional_layers"
    if (eq_ind and ca_support) or (ca_ind and eq_support):
        return "B_one_conditional_plus_query_support"
    return "C_other_association_evidence"


def compact_rows(frame: pd.DataFrame, columns: list[str]) -> str:
    columns = [c for c in columns if c in frame]
    return frame[columns].to_json(orient="records")


def complete_stage(out: Path, names: list[str]) -> None:
    """Only a fully written, hashed output set constitutes a completed stage."""
    outputs = {name: digest(out / name) for name in names}
    pending = out / "completion.json.tmp"
    pending.write_text(
        json.dumps(
            {
                "status": "complete",
                "stage": "manuscript_audit",
                "script_sha256": digest(Path(__file__)),
                "outputs": outputs,
            },
            indent=2,
        )
        + "\n"
    )
    pending.replace(out / "completion.json")


def receptor_relationships(root: Path, out: Path, chains: pd.DataFrame) -> None:
    summary_path = root / "results/rfu_state_validation_v1/observed_rfu_summary.tsv"
    if not summary_path.exists():
        return
    observed = read(summary_path).rename(
        columns={
            "rfu_id": "rfu_label",
            "v_call": "observed_dominant_TRBV",
            "fraction": "dominant_TRBV_fraction",
        }
    )
    observed.rfu_label = observed.rfu_label.astype(str)
    genes = (
        chains.loc[chains.layer == "eqtl"]
        .groupby("rfu_label", as_index=False)
        .agg(independent_egenes=("target", lambda x: json.dumps(sorted(set(x)))))
    )
    labels = chains[["rfu_label"]].drop_duplicates()
    relationships = labels.merge(genes, on="rfu_label", how="left", validate="one_to_one").merge(
        observed, on="rfu_label", how="left", validate="one_to_one"
    )
    relationships["rfuwas_number_compatibility"] = "not independently checksum-verified"
    relationships["interpretation"] = (
        "observed V calls; sparse single-cohort coverage, not reference-wide composition"
    )
    save(relationships, out / "rfu_trbv_regulatory_relationships.tsv")


def figures(root: Path, out: Path, counts: dict) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(7, 4), layout="constrained")
    names = [
        "eQTL represented",
        "eQTL query support",
        "eQTL conditional",
        "caQTL represented",
        "caQTL query support",
        "caQTL conditional",
        "Both conditional",
    ]
    values = [
        counts[layer][cat]["variants"]
        for layer in ("eqtl", "caqtl")
        for cat in ("tested", "lookup_bonferroni", "independent")
    ] + [counts["both_independent"]["variants"]]
    bars = ax.barh(names[::-1], values[::-1], color="#4C78A8")
    ax.bar_label(bars)
    ax.set(
        xlim=(0, 650),
        xlabel="Unique variants / 623 published RFU-QTL variants",
        title="Exact identity and conditional selection are separate evidence",
    )
    fig.savefig(out / "evidence_accounting.pdf")
    plt.close(fig)
    metrics_path = root / "results/rfu_state_validation_v1/donor_held_out_metrics.tsv"
    if metrics_path.exists():
        try:
            from .regulatory_rfu_state_validation import paired_deltas
        except ImportError:
            from regulatory_rfu_state_validation import paired_deltas

        metrics = read(metrics_path)
        paired = paired_deltas(metrics)
        fig, axes = plt.subplots(1, 2, figsize=(8, 3.5), layout="constrained")
        for ax, mode in zip(axes, ("clone", "cell"), strict=True):
            delta = paired.loc[paired.weighting == mode].set_index("held_out_donor").delta_log_loss
            ax.bar(delta.index, delta, color="#4C78A8")
            ax.axhline(0, color="black", linewidth=0.8)
            ax.set(
                title=f"{mode.capitalize()} weighted\nShared receptors excluded",
                ylabel="Added RFU log-loss change\npositive = worse",
            )
        fig.savefig(out / "rfu_beyond_trbv_heldout.pdf")
        plt.close(fig)


def workbook_audit(root: Path, variants: set[str]) -> dict:
    """Cache depends on workbook hash; source Table 5 is positive-hit-only GWAS coloc."""
    source = root / "sources/matos"
    workbook = source / "supplementary_tables.xlsx"
    if not workbook.exists():
        with zipfile.ZipFile(source / "PMC12870616_supplementary.zip") as archive:
            content = archive.read("media-2.xlsx")
        if hashlib.md5(content).hexdigest() != SUPPLEMENT_MD5:
            raise ValueError("Workbook checksum mismatch (not the supplement ZIP checksum)")
        workbook.write_bytes(content)
    if digest(workbook, "md5") != SUPPLEMENT_MD5:
        raise ValueError("Workbook checksum mismatch")
    cached = root / "prepared/matos_supplement_table5.tsv"
    manifest_path = root / "prepared/matos_supplement_manifest.json"
    query_hash = hashlib.sha256("\n".join(sorted(variants)).encode()).hexdigest()
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text())
        if (
            manifest["workbook_md5"] != SUPPLEMENT_MD5
            or manifest["table5_sha256"] != digest(cached)
            or manifest.get("query_variant_hash") != query_hash
        ):
            raise ValueError("Supplement cache identity mismatch")
        table = read(cached)
    else:
        import openpyxl

        book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        sheets = []
        source_ids = {"chr" + v.replace(":", "_") for v in variants}
        extracted = {}
        for sheet in book:
            first = list(sheet.iter_rows(min_row=1, max_row=3, values_only=True))
            sheets.append(
                {
                    "sheet": sheet.title,
                    "declared_rows": sheet.max_row,
                    "header_row": 3,
                    "headers": list(first[2]),
                    "title": list(first[0]),
                }
            )
            number = int(sheet.title.rsplit(" ", 1)[1])
            if number in (1, 3, 5, 6):
                retained, n_rows = [], 0
                for values in sheet.iter_rows(min_row=4, values_only=True):
                    row = {k: v for k, v in zip(first[2], values, strict=True) if k is not None}
                    if not any(v is not None for v in row.values()):
                        continue
                    n_rows += 1
                    ids = str(row.get("ChromBPNet_variant_ids", "")).split(";")
                    if (
                        number == 5
                        or row.get("variant_id") in source_ids
                        or bool(set(ids) & source_ids)
                    ):
                        retained.append(row)
                extracted[number] = pd.DataFrame(
                    retained, columns=[c for c in first[2] if c is not None]
                )
                sheets[-1]["nonempty_data_rows"] = n_rows
                sheets[-1]["retained_rows"] = len(retained)
                if number != 5:
                    save(
                        extracted[number],
                        root / f"prepared/matos_supplement_table{number}_exact.tsv",
                    )
        book.close()
        table = extracted[5]
        save(table, cached)
        manifest = {
            "workbook_md5": SUPPLEMENT_MD5,
            "workbook_bytes": workbook.stat().st_size,
            "table5_sha256": digest(cached),
            "table5_rows": len(table),
            "sheets": sheets,
            "query_variant_hash": query_hash,
        }
        manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")
    exact = table.loc[table.variant_id_GRC38.isin(variants)].copy()
    save(exact, root / "prepared/matos_source_gwas_coloc_exact.tsv")
    return {
        **manifest,
        "exact_rfu_variant_rows": len(exact),
        "table5_evidence": "QTL-GWAS coloc; not direct gene-peak PP.H4",
        "molecular_coloc_availability": "not in inspected workbook/summary archives",
        "absent_positive_hit_interpretation": "not a tested negative",
    }


def run(root: Path, out: Path) -> None:
    (out / "completion.json").write_text(
        json.dumps({"status": "running", "stage": "manuscript_audit"}) + "\n"
    )
    counts = json.loads((out / "overlap_counts.json").read_text())
    if counts["partial_caqtl"]:
        raise ValueError("Final dossiers require complete caQTL coverage")
    rfu = read(root / "prepared/rfuwas_data1_rfuqtl_grch38.tsv")
    if set(rfu.genome_build) != {"GRCh38"}:
        raise ValueError("GRCh38 required")
    raw = {k: read(out / f"matos_{k}_lookup_classified.tsv") for k in ("eqtl", "caqtl")}
    ind = {k: read(out / f"matos_{k}_independent_exact.tsv") for k in raw}
    context = read(root / "results/rfuqtl_cellstate_context.tsv").set_index("rfu_label")
    disease = read(root / "prepared/rfuwas_data4_annotations.tsv")
    rows, conditional = [], []
    for row in rfu.itertuples():
        v = row.variant_key
        e, c = (ind[k].loc[ind[k].variant_key == v] for k in ("eqtl", "caqtl"))
        en, cn = (
            raw[k].loc[(raw[k].variant_key == v) & raw[k].query_set_bonferroni_pass]
            for k in ("eqtl", "caqtl")
        )
        tier = evidence_tier(not e.empty, not c.empty, not en.empty, not cn.empty)
        if e.empty and c.empty and en.empty and cn.empty:
            continue
        ctx = (
            context.loc[str(row.rfu_label)].to_dict() if str(row.rfu_label) in context.index else {}
        )
        entry = {
            "variant": v,
            "rfu_label": str(row.rfu_label),
            "rfuqtl_beta": row.beta,
            "rfuqtl_pvalue": row.pvalue,
            "variant_location": location(row.chromosome, row.position),
            "independent_eqtl": not e.empty,
            "independent_caqtl": not c.empty,
            "source_molecular_colocalization": None,
            "PP_H4": None,
            "source_finemapping_support": None,
            "conditional_egenes": json.dumps(sorted(e.gene.unique())),
            "conditional_capeaks": json.dumps(sorted(c.peak.unique())),
            "eqtl_conditional_statistics": compact_rows(
                e, ["gene", "rank", "pvalue", "beta", "se", "source_signal_id", "source_file"]
            ),
            "caqtl_conditional_statistics": compact_rows(
                c, ["peak", "rank", "pvalue", "beta", "se", "source_signal_id", "source_file"]
            ),
            "query_supported_egenes": compact_rows(en, ["gene", "pvalue", "beta", "se"]),
            "query_supported_capeaks": compact_rows(cn, ["peak", "pvalue", "beta", "se"]),
            "downstream_rfuwas_phenotypes": compact_rows(
                disease.loc[disease.rfu_label == str(row.rfu_label)],
                ["phecode", "description", "pvalue", "effect", "group"],
            ),
            "evidence_tier": tier,
            "effect_orientation": "unknown",
            "notes": "Exact association/conditional selection; no RFU colocalization or mediation",
            **ctx,
        }
        rows.append(entry)
        for layer, sub in (("eqtl", e), ("caqtl", c)):
            for signal in sub.to_dict("records"):
                target = signal["gene" if layer == "eqtl" else "peak"]
                gene_class = (
                    (
                        "direct_TCR_gene"
                        if pd.Series([target]).str.match(r"^TR[ABDG][VJCD]").iloc[0]
                        else "non_TCR_gene"
                    )
                    if layer == "eqtl"
                    else None
                )
                anchor = int(signal["position"] - signal["start_distance"])
                conditional.append(
                    {
                        "variant": v,
                        "rfu_label": str(row.rfu_label),
                        "layer": layer,
                        "target": target,
                        "target_class": gene_class,
                        "source_signal_id": signal["source_signal_id"],
                        "conditional_rank": signal["rank"],
                        "conditional_pvalue": signal["pvalue"],
                        "conditional_beta": signal["beta"],
                        "conditional_se": signal["se"],
                        "target_anchor": anchor,
                        "anchor_location": location(signal["chromosome"], anchor),
                        "anchor_provenance": "variant position minus TensorQTL start_distance; not full interval",
                        "evidence_tier": tier,
                        "source_file": signal["source_file"],
                    }
                )
    candidates = pd.DataFrame(rows).sort_values(["evidence_tier", "variant", "rfu_label"])
    chains = pd.DataFrame(conditional).sort_values(
        ["layer", "source_signal_id", "variant", "rfu_label"]
    )
    high = candidates.loc[~candidates.evidence_tier.str.startswith("C")]
    save(candidates, out / "ranked_regulatory_candidates.tsv")
    save(high, out / "high_confidence_regulatory_candidates.tsv")
    save(chains, out / "high_confidence_independent_qtl_chains.tsv")
    save(read(out / "regulatory_hits.tsv"), out / "exact_regulatory_hits.tsv")
    summary = {
        "tier_counts": {
            tier: {
                "variant_rfu_pairs": len(sub),
                "variants": sub.variant.nunique(),
                "rfus": sub.rfu_label.nunique(),
            }
            for tier, sub in candidates.groupby("evidence_tier")
        },
        "conditional_target_classes": {
            str(cls): {
                "signals": sub.source_signal_id.nunique(),
                "targets": sub.target.nunique(),
                "variants": sub.variant.nunique(),
                "rfus": sub.rfu_label.nunique(),
            }
            for cls, sub in chains.groupby("target_class")
        },
        "source_signal_count_warning": "within-target conditional selections, not mutually LD-independent mechanisms",
        "molecular_coloc": workbook_audit(root, set(rfu.variant_key)),
    }
    # Compare annotations with all 59 RFU-QTL RFUs using explicit denominators.
    contexts = context.reset_index()
    annotation_summary = {}
    for group, labels in (("all_rfuqtl", set(rfu.rfu_label)), ("tier_A_B", set(high.rfu_label))):
        sub = contexts.loc[contexts.rfu_label.isin(labels)]
        annotation_summary[group] = {
            "rfus": len(labels),
            "CD4_CD8": sub.enrichment.fillna("unannotated").value_counts().to_dict(),
            "cellstate": sub.Enriched_Group.fillna("unannotated").value_counts().to_dict(),
        }
    summary["annotation_denominators"] = annotation_summary
    old = json.loads((root / "results/partial_caqtl_checkpoint/overlap_counts.json").read_text())
    summary["reconciliation"] = {
        "old_query_family": old["lookup_test_family_size"],
        "new_query_family": counts["lookup_test_family_size"],
        "eqtl_unchanged_rows": counts["eqtl"]["tested_variant_target_pairs"]
        == old["eqtl"]["tested_variant_target_pairs"],
        "added_caqtl_rows": counts["caqtl"]["tested_variant_target_pairs"]
        - old["caqtl"]["tested_variant_target_pairs"],
    }
    (out / "evidence_counts.json").write_text(json.dumps({**counts, **summary}, indent=2) + "\n")
    receptor_relationships(root, out, chains)
    figures(root, out, counts)
    import scrfu

    provenance = json.loads((out / "provenance.json").read_text())
    provenance["manuscript_audit"] = {
        "python": sys.executable,
        "python_version": platform.python_version(),
        "scrfu_import": scrfu.__file__,
        "git_sha": subprocess.check_output(["git", "rev-parse", "HEAD"], text=True).strip(),
        "script_sha256": digest(Path(__file__)),
        "inputs": {
            str(p): digest(p)
            for p in [
                root / "prepared/rfuwas_data1_rfuqtl_grch38.tsv",
                *[
                    out / f"matos_{k}_{suffix}.tsv"
                    for k in raw
                    for suffix in ("lookup_classified", "independent_exact")
                ],
            ]
        },
        "TRB_reference": {
            "NCBI_Gene": 6957,
            "accession": "NC_000007.14",
            "inclusive_interval": TRB,
        },
        "matos_code_commit": "4b50d55a9bf8339f3625a4de6da9bf618c288b48",
    }
    (out / "provenance.json").write_text(json.dumps(provenance, indent=2) + "\n")
    names = [
        "evidence_counts.json",
        "ranked_regulatory_candidates.tsv",
        "high_confidence_regulatory_candidates.tsv",
        "high_confidence_independent_qtl_chains.tsv",
        "exact_regulatory_hits.tsv",
        "evidence_accounting.pdf",
        "provenance.json",
    ]
    if (root / "results/rfu_state_validation_v1/observed_rfu_summary.tsv").exists():
        names.append("rfu_trbv_regulatory_relationships.tsv")
    if (root / "results/rfu_state_validation_v1/donor_held_out_metrics.tsv").exists():
        names.append("rfu_beyond_trbv_heldout.pdf")
    complete_stage(out, names)
    print(json.dumps({k: v for k, v in summary.items() if k != "molecular_coloc"}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    run(args.root, args.out)
