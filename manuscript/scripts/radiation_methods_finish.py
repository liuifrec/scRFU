"""Assemble manuscript tables/figures from completed analyses, without refitting.

This is a reporting stage. It verifies all prior completion manifests, audits
the small author-resource set, and keeps datasets/intervals as separate rows.
"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

from scrfu.io import file_sha256

from ._figure_common import clean_axis, panel_label, save_figure, set_style
from .gse280982_analysis import finish, identity, start, verified_stage
from .radiation_methods import json_save, save


def read(root: Path, name: str) -> pd.DataFrame:
    return pd.read_csv(root / name, sep="\t")


def external_results(root: Path) -> pd.DataFrame:
    keys = ["donor", "compartment", "interval"]
    all_d = read(root, "multiscale_pairs.tsv")
    primary = all_d[
        all_d.policy.eq("threshold") & all_d.weighting.eq("cell") & all_d.grouping.eq("RFU")
    ].set_index(keys)
    columns = [
        "visit_before",
        "visit_after",
        "before_qualified_cells",
        "after_qualified_cells",
        "d_clone",
        "d_group",
        "aggregation_cancellation",
    ]
    result = primary[columns].copy()
    for name, subset in [
        (
            "unique_receptor_RFU_TV",
            all_d[
                all_d.policy.eq("threshold")
                & all_d.weighting.eq("unique_clone")
                & all_d.grouping.eq("RFU")
            ],
        ),
        (
            "TRBV_TV",
            all_d[
                all_d.policy.eq("threshold")
                & all_d.weighting.eq("cell")
                & all_d.grouping.eq("TRBV")
            ],
        ),
        (
            "VJ_TV",
            all_d[
                all_d.policy.eq("threshold")
                & all_d.weighting.eq("cell")
                & all_d.grouping.eq("TRBV_TRBJ")
            ],
        ),
    ]:
        result[name] = subset.set_index(keys).d_group
    depth = read(root, "cell_depth_sensitivity.tsv").groupby(keys)
    result["depth_cells_each_visit"] = depth.cells_per_visit.first()
    for suffix, q in [("q025", 0.025), ("median", 0.5), ("q975", 0.975)]:
        result["depth_RFU_TV_" + suffix] = depth.d_group.quantile(q)
    control = read(root, "fixed_group_controls.tsv").groupby(keys)
    result["matched_map_cancellation_median"] = control.aggregation_cancellation.median()
    result["matched_map_fraction_labels_changed"] = control.pair_labels_changed_fraction.median()
    dom = read(root, "clone_dominance_sensitivity.tsv").set_index(keys)
    result["dominant_removed_RFU_TV"] = dom.d_group
    result = result.join(
        dom[
            [
                "dominant_fraction_before",
                "dominant_fraction_after",
                "retained_fraction_before",
                "retained_fraction_after",
            ]
        ]
    )
    sim = read(root, "similarity.tsv")
    result["RFU_cosine"] = (
        sim[sim.policy.eq("threshold") & sim.weighting.eq("cell")].set_index(keys).rfu_cosine
    )
    persistence = read(root, "persistence_summary.tsv")
    for detection in (1, 5):
        sub = persistence[persistence.detection_min_cells.eq(detection)].set_index(keys)
        for col in [
            "detected_union_RFUs",
            "persistent_RFUs",
            "persistent_over_union",
            "persistent_without_shared_receptor_fraction",
        ]:
            result[f"{col}_min{detection}cells"] = sub[col]
    return result.reset_index()


def cross_application(rp: Path, dev: Path, external: Path) -> pd.DataFrame:
    result = []
    d = read(rp, "rp1_14_multiscale_pairs.tsv")
    d = d[
        d.primary_pair & d.policy.eq("threshold") & d.weighting.eq("reads") & d.grouping.eq("RFU")
    ]
    coverage = read(rp, "rp1_14_sample_coverage.tsv")
    p = read(rp, "rp1_14_persistence_donor_summary.tsv")
    depth = read(rp, "rp1_14_empirical_read_sensitivity.tsv")
    for comp, part in d.groupby("compartment"):
        c = coverage[coverage.compartment.eq(comp)]
        result.append(
            {
                "dataset": "RP1-14",
                "compartment": comp,
                "participants": part.donor.nunique(),
                "interval": "earliest_to_latest_14.95_to_24.86_years",
                "available_visits": len(c),
                "sampling_unit": "sequencing_read",
                "depth_min": int(c.all_source_reads.min()),
                "depth_max": int(c.all_source_reads.max()),
                "coverage_denominator": "all_source_reads",
                "coverage_min": c.fixed_primary_fraction_all_reads.min(),
                "coverage_max": c.fixed_primary_fraction_all_reads.max(),
                "receptor_TV_median": part.d_clone.median(),
                "RFU_TV_median": part.d_group.median(),
                "cancellation_median": part.aggregation_cancellation.median(),
                "persistent_union_median": p[
                    p.compartment.eq(comp)
                ].persistent_fraction_union.median(),
                "persistence_detection": "5_reads_each_visit",
                "depth_RFU_TV_median": depth[depth.compartment.eq(comp)]
                .groupby("donor")
                .d_group.median()
                .median(),
                "depth_scheme": "50_without_replacement_read_draws_500_reads",
                "limitation": "six_same_people_across_compartments; incomplete_reused_map; read_not_template; no_radiation_exposure_comparison",
            }
        )
    c = read(dev, "gse190905_paired_manifest.tsv")
    d = read(dev, "gse190905_multiscale_pairs.tsv")
    d = d[
        d.subset.eq("all_T")
        & d.policy.eq("threshold")
        & d.weighting.eq("cell")
        & d.grouping.eq("RFU")
    ]
    p = read(dev, "gse190905_rfu_observed_persistence.tsv")
    p = p[p.subset.eq("all_T") & p.policy.eq("threshold") & p.detection_min_cells.eq(1)]
    persistence = []
    for _, group in p.groupby("donor"):
        persistence.append(
            (group.observed_before & group.observed_after).sum()
            / (group.observed_before | group.observed_after).sum()
        )
    depth = read(dev, "gse190905_empirical_depth_sensitivity.tsv")
    result.append(
        {
            "dataset": "GSE190905",
            "compartment": "blood_primary_TRB",
            "participants": 6,
            "available_visits": 12,
            "interval": "pre_to_post_SABR",
            "sampling_unit": "primary_TRB_cell",
            "depth_min": int(c.tcr_cells.min()),
            "depth_max": int(c.tcr_cells.max()),
            "coverage_denominator": "primary_TRB_cells",
            "coverage_min": c.threshold_coverage_of_tcr.min(),
            "coverage_max": c.threshold_coverage_of_tcr.max(),
            "receptor_TV_median": d.d_clone.median(),
            "RFU_TV_median": d.d_group.median(),
            "cancellation_median": d.aggregation_cancellation.median(),
            "persistent_union_median": float(np.median(persistence)),
            "persistence_detection": "1_cell_each_visit",
            "depth_RFU_TV_median": depth.groupby("donor").d_group.median().median(),
            "depth_scheme": "50_without_replacement_cell_draws_cap500_one_pair404",
            "limitation": "six_donors; four_SABR_two_prior_systemic; source_demultiplexing; final_seven_donor_paper_not_fully_released_in_cache",
        }
    )
    e = external_results(external)
    c = read(external, "visit_support.tsv")
    for (comp, interval), part in e.groupby(["compartment", "interval"]):
        visits = set(part.visit_before) | set(part.visit_after)
        sub = c[c.compartment.eq(comp) & c.donor.isin(part.donor) & c.visit.isin(visits)]
        result.append(
            {
                "dataset": "GSE280982",
                "compartment": comp,
                "participants": part.donor.nunique(),
                "available_visits": len(sub),
                "interval": interval,
                "sampling_unit": "GEX_matched_primary_TRB_cell",
                "depth_min": int(sub.primary_TRB_cells.min()),
                "depth_max": int(sub.primary_TRB_cells.max()),
                "coverage_denominator": "GEX_matched_primary_TRB_cells",
                "coverage_min": sub.coverage.min(),
                "coverage_max": sub.coverage.max(),
                "receptor_TV_median": part.d_clone.median(),
                "RFU_TV_median": part.d_group.median(),
                "cancellation_median": part.aggregation_cancellation.median(),
                "persistent_union_median": part.persistent_over_union_min1cells.median(),
                "persistence_detection": "1_cell_each_visit",
                "depth_RFU_TV_median": part.depth_RFU_TV_median.median(),
                "depth_scheme": "50_without_replacement_cell_draws_min_pair_depth_cap500",
                "limitation": "same_donors_recur_across_intervals; radiation_day_small; one_tumor_final_visit_missing; no_direct_author_cell_states; no_pooled_effect",
            }
        )
    return pd.DataFrame(result)


def metadata_audit(workspace: Path, out: Path) -> None:
    root = workspace / "sources/GSE280982"
    code = root / "author_analysis.R"
    if hashlib.md5(code.read_bytes()).hexdigest() != "d3f76420491638247349bf0ef13e956d":
        raise ValueError("Published author R-code MD5 mismatch.")
    book = root / "41467_2025_60827_MOESM4_ESM.xlsx"
    wb = openpyxl.load_workbook(book, read_only=True, data_only=True)
    save(
        pd.DataFrame([{"sheet": s.title, "rows": s.max_row, "columns": s.max_column} for s in wb]),
        out,
        "metadata_workbook_inventory.tsv",
    )
    wb.close()
    json_save(
        {
            "status": "no_directly_reusable_author_barcode_to_state_table_found_in_bounded_check",
            "checked": [
                "cached_GEO_series_and_processed_inventory",
                "paper_data_and_code_availability",
                "Zenodo_15258024_file_inventory",
                "Zenodo_15265351_file_inventory",
                "25_source_workbook_sheets",
                "author_R_code",
            ],
            "paper": "https://doi.org/10.1038/s41467-025-60827-w",
            "code_source": "https://zenodo.org/records/15258024",
            "code_published_md5": "d3f76420491638247349bf0ef13e956d",
            "workbook_sha256": file_sha256(book),
            "code_sha256": file_sha256(code),
            "workbook_scope": "DEGs, aggregate cluster proportions, selected TProlif_Tox clonotypes, per-sample receptor lists, binding/experimental summaries; Fig8 barcode/CDR3 lacks state",
            "code_scope": "author cluster-to-state labels applied after RNA processing; no released barcode-to-cluster object identified",
            "decision": "external_state_analysis_unavailable; no_clustering_or_inference_from_selected_receptors",
            "limitation": "bounded availability audit, not proof that no author-held metadata exists",
        },
        out / "metadata_availability.json",
    )


def figures(workspace: Path, out: Path, e: pd.DataFrame) -> None:
    set_style()
    coverage = read(workspace / "results/gse280982_v1", "visit_support.tsv")
    fig, axes = plt.subplots(2, 2, figsize=(10, 7), constrained_layout=True)
    ax = axes[0, 0]
    for (donor, comp), rows in coverage.groupby(["donor", "compartment"]):
        if rows.qualified_cells.notna().any():
            ax.plot(
                rows.visit,
                rows.qualified_cells,
                "o-" if comp == "tumor" else "s--",
                label=f"{donor} {comp}",
            )
    ax.axhline(100, color="black", linestyle=":")
    ax.set(
        xticks=[1, 2, 3],
        xticklabels=["Pre", "Last RT day", "~6 weeks"],
        ylabel="Qualified primary-TRB cells",
        title="Support ≥100; unavailable visits are gaps",
    )
    ax.legend(ncol=2)
    labels = [
        f"{r.donor} {r.compartment}\n{int(r.visit_before)}→{int(r.visit_after)}"
        for r in e.itertuples()
    ]
    x = np.arange(len(e))
    ax = axes[0, 1]
    ax.plot(x, e.d_clone, "o", label="Receptor TV")
    ax.plot(x, e.d_group, "s", label="RFU TV")
    for i, row in e.iterrows():
        ax.plot([i, i], [row.d_group, row.d_clone], color="gray", linewidth=1)
    ax.set(ylabel="Total variation", ylim=(0, 1.05), title="Each donor and interval retained")
    ax.legend()
    ax = axes[1, 0]
    ax.plot(x, e.persistent_over_union_min1cells, "o", label="Persistent / detected union")
    ax.plot(
        x,
        e.persistent_without_shared_receptor_fraction_min1cells,
        "s",
        label="No shared receptor / persistent",
    )
    ax.set(
        ylabel="Fraction (RFUs detected at ≥1 cell)",
        ylim=(0, 1.05),
        title="Distinct persistence denominators",
    )
    ax.legend()
    ax = axes[1, 1]
    ax.plot(x, e.aggregation_cancellation, "o", label="Observed")
    ax.plot(x, e.matched_map_cancellation_median, "s", label="Matched-map median")
    # Display depth TV separately in the source table; do not imply cancellation
    # equals the difference of independently summarized medians.
    ax.set(ylabel="Aggregation cancellation", title="Matched grouping controls")
    ax.legend()
    for ax in [axes[0, 1], axes[1, 0], axes[1, 1]]:
        ax.set(xticks=x, xticklabels=labels)
        ax.tick_params(axis="x", labelrotation=50)
    for letter, ax in zip("ABCD", axes.flat, strict=True):
        panel_label(ax, letter)
        clean_axis(ax)
    save_figure(fig, out, "gse280982_external_final")
    plt.close(fig)
    # A table, rather than a decorative figure, documents the repair test.
    q = read(workspace / "results/rp1_14_repair_qc_v1", "parity_summary.tsv")
    q = q[~q.factor.isin(["donor", "visit"])]
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.axis("off")
    labels = ["Stratum", "Unique AA", "Label errors", "Threshold errors", "Score errors >10⁻¹²"]
    table = ax.table(
        cellText=q[
            [
                "level",
                "n_unique_receptors",
                "label_mismatches",
                "threshold_mismatches",
                "score_mismatches_at_1e_12",
            ]
        ].values,
        colLabels=labels,
        loc="center",
        colWidths=[0.29, 0.14, 0.16, 0.18, 0.23],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(8)
    table.scale(1, 1.5)
    ax.set_title(
        "RP1-14 repair QC: 1,600 unique receptors; strata reuse the same assay\n330,433 changed / 29,097 unchanged AA labels; all historical RFU/score vectors identical",
        fontsize=10,
        pad=12,
    )
    fig.text(
        0.05,
        0.025,
        "RFU conflicts: 41,572 → 0 AA groups. Score conflicts: 41,581 → 0. Bounded parity does not prove the missing historical execution manifest.",
        fontsize=7,
    )
    save_figure(fig, out, "rp1_14_repair_qc_table")
    plt.close(fig)
    schema = pd.DataFrame(
        [
            {
                "step": "Inputs and design",
                "content": "Donor × visit × compartment\nCounts and sampling unit\nMissing visits retained",
            },
            {
                "step": "Fixed reference",
                "content": "Published RFU reference hash\nThreshold and coverage\nMap fixed across visits",
            },
            {
                "step": "Separate measurements",
                "content": "Receptor / group TV\nAggregation cancellation\nReceptor turnover within RFUs",
            },
            {
                "step": "Interpretation checks",
                "content": "Depth, clone and matched maps\nSource cellular annotations\nTrace external evidence",
            },
        ]
    )
    save(schema, out, "framework_schema.tsv")
    fig, ax = plt.subplots(figsize=(10, 3))
    ax.set(xlim=(0, 4), ylim=(0, 1))
    ax.axis("off")
    for i, row in schema.iterrows():
        ax.text(i + 0.5, 0.62, row.step, ha="center", va="center", fontsize=10, weight="bold")
        ax.text(
            i + 0.5,
            0.38,
            row.content,
            ha="center",
            va="center",
            fontsize=7,
            linespacing=1.7,
            bbox={"boxstyle": "round,pad=.6", "facecolor": "#e8f2f8", "edgecolor": "#0072B2"},
        )
    fig.text(
        0.5,
        0.08,
        "Same definitions across designs; no pooling of bulk reads, cells or biological effects",
        ha="center",
        fontsize=9,
    )
    save_figure(fig, out, "framework")
    plt.close(fig)


def run(workspace: Path, output_name: str = "manuscript_finish_v1_1") -> None:
    dirs = {
        name: workspace / "results" / name
        for name in ["rp1_14_v1", "development_v1_1", "gse280982_v1", "rp1_14_repair_qc_v1"]
    }
    for root in dirs.values():
        verified_stage(root)
    inputs = {name: p / "completion.json" for name, p in dirs.items()}
    inputs.update(
        {
            name: workspace / "sources/GSE280982" / name
            for name in ["author_analysis.R", "41467_2025_60827_MOESM4_ESM.xlsx"]
        }
    )
    scientific = identity(
        inputs, (run, read, external_results, cross_application, metadata_audit, figures)
    )
    out = workspace / "results" / output_name
    fingerprint = start(out, scientific)
    if fingerprint is None:
        return
    e = external_results(dirs["gse280982_v1"])
    save(e, out, "external_donor_results.tsv")
    summary = cross_application(dirs["rp1_14_v1"], dirs["development_v1_1"], dirs["gse280982_v1"])
    save(summary, out, "cross_application_summary.tsv")
    metadata_audit(workspace, out)
    figures(workspace, out, e)
    rows = []
    for figure, panels, source in [
        ("framework.pdf", "schema", out / "framework_schema.tsv"),
        ("gse280982_external_final.pdf", "A", dirs["gse280982_v1"] / "visit_support.tsv"),
        ("gse280982_external_final.pdf", "B,C,D", out / "external_donor_results.tsv"),
        ("rp1_14_repair_qc_table.pdf", "table", dirs["rp1_14_repair_qc_v1"] / "parity_summary.tsv"),
    ]:
        rows.append(
            {
                "figure": figure,
                "panels": panels,
                "source_table": str(source),
                "source_sha256": file_sha256(source),
                "configuration_and_code": "provenance.json links verified parent manifests",
                "completion": "completion.json",
                "unit": "donor trajectory for external analysis; one receptor assay for repair QC; framework is a schema",
                "limits": "no population radiation effect, antigen persistence, RFU-specific stability or full historical execution identity",
                "command": "python -m manuscript.scripts.radiation_methods_finish --workspace WORKSPACE",
            }
        )
    save(pd.DataFrame(rows), out, "figure_source_index.tsv")
    save(
        pd.DataFrame(
            [{"source_table": p.name, "sha256": file_sha256(p)} for p in sorted(out.glob("*.tsv"))]
        ),
        out,
        "source_table_index.tsv",
    )
    json_save(
        {
            "status": "all_scoped_reporting_outputs_generated",
            "cross_application_rows": len(summary),
            "external_donor_intervals": len(e),
            "no_biological_measurement_rerun": True,
            "independent_external_people": 3,
            "paired_external_blood_people": 1,
            "RP1_14_primary_results_unchanged": True,
            "prediction_and_regulatory_results_unchanged": True,
        },
        out / "evidence_counts.json",
    )
    finish(out, fingerprint)
    print("Manuscript assembly complete; prior endpoints reused.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace", type=Path, required=True)
    run(parser.parse_args().workspace)
